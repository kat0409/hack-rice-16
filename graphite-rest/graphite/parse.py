"""Turn a notes file into structured text blocks (design-doc.md §8.3).

A Block is the smallest unit the parser is confident about: a paragraph, plus
where it came from. Chunking (chunk.py) decides how blocks get grouped; parsing
only extracts and normalizes, and never merges across a structural boundary.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}


class ParseError(Exception):
    """Carries a stable machine-readable code for the API error envelope (§11.6)."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class Block:
    text: str
    page: int | None = None
    section_path: str | None = None


@dataclass
class ParsedDocument:
    blocks: list[Block]
    page_count: int | None
    used_ocr: bool = False


def normalize(text: str) -> str:
    """Normalize without destroying content.

    §8.3 warns against aggressive cleanup: repeated text may be a definition or a
    formula, so this only fixes encoding and layout artifacts, and deliberately
    does not strip headers, footers, or repeated lines.
    """
    text = unicodedata.normalize("NFKC", text)
    # Words split across a line break by PDF/print hyphenation: "recur-\nsion".
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Collapse horizontal whitespace only; newlines carry paragraph structure.
    text = re.sub(r"[ \t]+", " ", text)
    # 3+ newlines is always a formatting artifact, not meaning.
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def _parse_pdf(path: Path, ocr: bool = False) -> ParsedDocument:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise ParseError("PARSE_FAILED", f"Could not open PDF: {exc}") from exc

    if doc.needs_pass:
        raise ParseError(
            "PARSE_ENCRYPTED",
            "This PDF is password-protected. Remove the password and re-upload.",
        )

    blocks: list[Block] = []
    for page_number, page in enumerate(doc, start=1):
        for paragraph in _split_paragraphs(normalize(page.get_text())):
            blocks.append(Block(text=paragraph, page=page_number))

    page_count = doc.page_count
    doc.close()

    used_ocr = False
    if not blocks and ocr:
        # No text layer: a scan or a photographed page. Transcribing it sends
        # page images to the model provider, which is why this never happens
        # unless the caller explicitly asked for it (see ocr.py).
        from graphite.ocr import ocr_pdf

        blocks = ocr_pdf(path)
        used_ocr = True

    if not blocks:
        # §20.5 PARSE_EMPTY: almost always a scanned PDF. Never send empty content
        # to the model — fail with something the student can act on.
        raise ParseError(
            "PARSE_EMPTY",
            "We could not find selectable text in this file. It is likely a scan "
            "or photographed handwriting. Re-upload with OCR enabled to have it "
            "transcribed, or export it as a text-based PDF, DOCX, Markdown, or TXT.",
        )
    return ParsedDocument(blocks=blocks, page_count=page_count, used_ocr=used_ocr)


def _parse_docx(path: Path) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError("PARSE_FAILED", f"Could not open DOCX: {exc}") from exc

    blocks: list[Block] = []
    heading_stack: list[str] = []

    for paragraph in document.paragraphs:
        text = normalize(paragraph.text)
        if not text:
            continue

        level = _docx_heading_level(paragraph)
        if level is not None:
            del heading_stack[level - 1 :]
            heading_stack.append(text)
            continue

        blocks.append(
            Block(text=text, section_path=" > ".join(heading_stack) or None)
        )

    if not blocks:
        raise ParseError("PARSE_EMPTY", "This document contains no readable text.")
    return ParsedDocument(blocks=blocks, page_count=None)


def _docx_heading_level(paragraph) -> int | None:
    name = (paragraph.style.name or "") if paragraph.style else ""
    match = re.fullmatch(r"Heading (\d+)", name)
    return int(match.group(1)) if match else None


def _parse_markdown(path: Path) -> ParsedDocument:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ParseError("PARSE_FAILED", "File is not valid UTF-8 text.") from exc

    blocks: list[Block] = []
    heading_stack: list[str] = []
    buffer: list[str] = []
    in_code_fence = False

    def flush() -> None:
        if not buffer:
            return
        body = normalize("\n".join(buffer))
        for paragraph in _split_paragraphs(body):
            blocks.append(
                Block(text=paragraph, section_path=" > ".join(heading_stack) or None)
            )
        buffer.clear()

    for line in raw.splitlines():
        if line.lstrip().startswith("```"):
            in_code_fence = not in_code_fence
            buffer.append(line)
            continue

        # A '#' inside a fenced code block is code, not a heading.
        heading = None if in_code_fence else re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            del heading_stack[level - 1 :]
            heading_stack.append(heading.group(2).strip())
        else:
            buffer.append(line)
    flush()

    if not blocks:
        raise ParseError("PARSE_EMPTY", "This file contains no readable text.")
    return ParsedDocument(blocks=blocks, page_count=None)


_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".md": _parse_markdown,
    ".markdown": _parse_markdown,
    ".txt": _parse_markdown,
}


def parse_file(path: Path, *, ocr: bool = False) -> ParsedDocument:
    """Parse a notes file into blocks.

    `ocr` only affects PDFs with no text layer, and defaults to False because
    transcription sends page images off the machine (§8.11).
    """
    suffix = path.suffix.lower()
    if suffix not in _PARSERS:
        raise ParseError(
            "UNSUPPORTED_FILE_TYPE",
            f"{suffix or 'This file type'} is not supported. "
            f"Use one of: {', '.join(sorted(SUPPORTED_SUFFIXES))}.",
        )
    if suffix == ".pdf":
        return _parse_pdf(path, ocr=ocr)
    return _PARSERS[suffix](path)
